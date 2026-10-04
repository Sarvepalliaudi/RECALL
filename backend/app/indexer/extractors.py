"""
File Content Extractors for RECALL.
Extracts structured text from PDF, DOCX, TXT, MD, PY, JAVA, JS, TS, HTML, CSV,
and prepares images for OCR.
"""
import hashlib
import io
from pathlib import Path
from typing import Dict, Any, Optional

SUPPORTED_TEXT_EXTENSIONS = {
    ".txt", ".md", ".py", ".java", ".js", ".ts", ".tsx", ".jsx",
    ".html", ".css", ".csv", ".json", ".xml", ".yaml", ".yml", ".sh", ".sql"
}
SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}
SUPPORTED_DOC_EXTENSIONS = {".pdf", ".docx"}

ALL_SUPPORTED_EXTENSIONS = (
    SUPPORTED_TEXT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS | SUPPORTED_DOC_EXTENSIONS
)


def compute_sha256(data: bytes) -> str:
    """Calculate SHA-256 hash of byte content."""
    return hashlib.sha256(data).hexdigest()


def extract_text_from_pdf(data: bytes) -> str:
    """Extract plain text from PDF bytes using PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=data, filetype="pdf")
        pages_text = []
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                pages_text.append(f"[Page {page_num + 1}]\n{text.strip()}")
        return "\n\n".join(pages_text)
    except Exception as e:
        return f"[Error extracting PDF: {str(e)}]"


def extract_text_from_docx(data: bytes) -> str:
    """Extract plain text from DOCX bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except Exception as e:
        return f"[Error extracting DOCX: {str(e)}]"


def extract_plain_text(data: bytes) -> str:
    """Safely decode plain text or code files using utf-8 with fallback encodings."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_content_from_file(filename: str, data: bytes) -> Dict[str, Any]:
    """
    Dispatcher to extract content, hash, and metadata from arbitrary file bytes.
    Never raises an uncaught exception; returns structured diagnostic status.
    """
    ext = Path(filename).suffix.lower()
    content_hash = compute_sha256(data)
    size_bytes = len(data)

    result = {
        "filename": filename,
        "extension": ext,
        "size_bytes": size_bytes,
        "content_hash": content_hash,
        "text": "",
        "is_image": False,
        "error": None,
    }

    if ext not in ALL_SUPPORTED_EXTENSIONS:
        result["error"] = f"Unsupported file extension: {ext}"
        return result

    try:
        if ext in SUPPORTED_TEXT_EXTENSIONS:
            result["text"] = extract_plain_text(data)
        elif ext == ".pdf":
            result["text"] = extract_text_from_pdf(data)
        elif ext == ".docx":
            result["text"] = extract_text_from_docx(data)
        elif ext in SUPPORTED_IMAGE_EXTENSIONS:
            result["is_image"] = True
            # Image OCR will be invoked if OCR is enabled
    except Exception as e:
        result["error"] = str(e)

    return result
